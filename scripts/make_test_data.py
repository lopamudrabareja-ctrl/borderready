"""Fills the judges' template (from row 4) with 17 test bookings, keeping its sheets and formatting.
Usage: python3 scripts/make_test_data.py <template.xlsx> <output.xlsx>"""
import sys
from openpyxl import load_workbook

UK, CN, IN, VN, TR, DE = "United Kingdom", "China", "India", "Vietnam", "Turkey", "Germany"

# Booking ID, Seller, Origin city, Origin country, Dest city, Dest country, Dest type, Ready date, Priority,
# Pallets, Cartons, Declared value, Currency, Gross kg, Appointment
BOOKINGS = [
    ["T-01", "Leeds Linens Ltd", "Leeds", UK, "Paris", "France", "Business address", "16-11-2026", "Fastest", 2, 40, 2400, "GBP", 500, "No"],
    ["T-02", "Hearth & Home", "London", UK, "Amsterdam", "Netherlands", "Warehouse", "18-11-2026", "Cheapest", 2, 80, 4800, "GBP", 700, "Yes"],
    ["T-03", "Brum Leather Co", "Birmingham", UK, "Berlin", "Germany", "Business address", "20-11-2026", "Cheapest", 1, 20, 2800, "GBP", 300, "No"],
    ["T-04", "Pennine Pottery", "Manchester", UK, "Lyon", "France", "Business address", "17-11-2026", "Balanced", 2, 80, 3000, "GBP", 600, "No"],
    ["T-05", "Pennine Pottery", "Manchester", UK, "Lyon", "France", "Business address", "17-11-2026", "Balanced", 2, 80, 3380, "GBP", 600, "No"],
    ["T-06", "Clearview Glass", "London", UK, "Hamburg", "Germany", "Business address", "19-11-2026", "Cheapest", 1, 20, 1000, "GBP", 300, "No"],
    ["T-07", "Dales Knitwear", "Leeds", UK, "Rotterdam", "Netherlands", "Business address", "21-11-2026", "Most reliable", 1, 60, 3500, "GBP", 350, "No"],
    ["T-08", "Ironworks Kitchen", "Birmingham", UK, "Frankfurt", "Germany", "Warehouse", "23-11-2026", "Cheapest", 2, 100, 6000, "GBP", 1150, "Yes"],
    ["T-09", "Stone & Garden", "London", UK, "Paris", "France", "Business address", "24-11-2026", "Fastest", 1, 20, 800, "GBP", 800, "No"],
    ["T-10", "Gadget Hub UK", "Manchester", UK, "Amsterdam", "Netherlands", "Business address", "18-11-2026", "Cheapest", 1, 15, 2100, "GBP", 250, "No"],
    ["T-11", "Little Oak Toys", "Leeds", UK, "Berlin", "Germany", "Business address", "20-11-2026", "Balanced", 1, 20, 2700, "GBP", 200, "No"],
    ["T-12", "Hartley Books", "London", UK, "Lyon", "France", "Business address", "22-11-2026", "Cheapest", 1, 20, 2700, "GBP", 200, "No"],
    ["T-13", "Soft Cotton Co", "Manchester", UK, "Hamburg", "Germany", "Warehouse", "25-11-2026", "Cheapest", 2, 40, 3200, "GBP", 600, "No"],
    ["T-14", "Bright Lamps", "Birmingham", UK, "Rotterdam", "Netherlands", "Business address", "26-11-2026", "Cheapest", 1, 10, 1800, "GBP", 250, "No"],
    ["T-15", "Leeds Linens Ltd", "Leeds", UK, "Paris", "France", "Business address", "10-10-2026", "Cheapest", 1, 10, 1000, "GBP", 200, "No"],
    ["T-16", "Velvet Interiors", "London", UK, "Frankfurt", "Germany", "Warehouse", "27-11-2026", "Most reliable", 1, 20, 4500, "GBP", 380, "No"],
    ["T-17", "Two Leaves Tea", "London", UK, "Amsterdam", "Netherlands", "Business address", "19-11-2026", "Cheapest", 1, 10, 1200, "GBP", 100, "No"],
]

# Booking ID, SKU, Description, HS code, Quantity, Unit price, Units per carton, Cartons, Unit kg, Country of origin
ITEMS = [
    ["T-01", "LIN-100", "Cotton bed linen set", 630231, 400, 6, 10, 40, 1, IN],
    ["T-02", "CDL-200", "Scented soy candle", 340600, 600, 4, 12, 50, 0.5, CN],
    ["T-02", "VAS-210", "Ceramic flower vase", 691390, 240, 10, 8, 30, 1.2, VN],
    ["T-03", "WAL-300", "Leather wallet", None, 200, 8, 20, 10, 0.4, TR],
    ["T-03", "BLT-310", "Leather belt", 42033, 100, 12, 10, 10, 0.6, TR],
    ["T-04", "PLT-400", "Ceramic dinner plate", 691110, 960, 3.5, 12, 80, 0.5, UK],
    ["T-05", "PLT-400", "Ceramic dinner plate", 691110, 960, 3.5, 12, 80, 0.5, UK],
    ["T-06", "TUM-600", "Glass tumbler, set of 4", 701337, 500, 2, 24, 20, 0.3, CN],
    ["T-07", "SCF-700", "Wool scarf", 611710, 500, 7, 10, 50, 0.2, UK],
    ["T-08", "PAN-800", "Cast iron frying pan", 732391, 400, 15, 4, 100, 2.5, CN],
    ["T-09", "PAV-900", "Garden paving stone", 680221, 40, 20, 2, 20, 18, IN],
    ["T-10", "PWB-1000", "Power bank 20000mAh", 850760, 100, 15, 20, 5, 0.4, CN],
    ["T-10", "HSP-1010", "Hair spray aerosol 300 ml", 330530, 120, 3, 24, 5, 0.35, UK],
    ["T-10", "NPR-1020", "Nail polish remover (acetone)", 330430, 120, 2, 24, 5, 0.2, UK],
    ["T-11", "DEO-1100", "Lithium-free natural deodorant stick", 330720, 300, 5, 30, 10, 0.1, UK],
    ["T-11", "TRN-1110", "Battery-free wooden toy train", 950300, 100, 12, 10, 10, 0.8, DE],
    ["T-12", "BOK-1200", "Hardback cookbook", 490199, 300, 9, 15, 20, 1, UK],
    ["T-13", "TWL-1300", "Cotton bath towel", 630260, 800, 4, 20, 40, 0.6, IN],
    ["T-14", "LMP-1400", "Brass desk lamp", 940520, 100, 18, 10, 10, 1.5, None],
    ["T-15", "TTW-1500", "Linen tea towel", 630291, 500, 2, 50, 10, 0.1, UK],
    ["T-16", "CUS-1600", "Silk cushion cover", None, 200, 10, 20, 10, 0.3, CN],
    ["T-16", "THR-1610", "Velvet throw blanket", 630140, 100, 25, 10, 10, 1.2, None],
    ["T-17", "TEA-1700", "Loose leaf green tea 250 g", "0902300", 200, 6, 20, 10, 0.3, IN],
]

def fill(ws, rows, start=4):
    for r, row in enumerate(rows, start=start):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)

template, out = sys.argv[1], sys.argv[2]
wb = load_workbook(template)
fill(wb["Bookings"], BOOKINGS)
fill(wb["Items"], ITEMS)
wb.save(out)
print(f"Wrote {out}: {len(BOOKINGS)} bookings, {len(ITEMS)} items (template example row kept)")
